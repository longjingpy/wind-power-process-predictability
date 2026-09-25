"""Build current Applied Energy cover letter and highlights DOCX files."""
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "manuscript/applied_energy"


def style_document(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.5)
    for name in ("Title", "Heading 1", "Heading 2"):
        style = doc.styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.color.rgb = None


def add_para(doc, text="", bold=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    run = p.add_run(text)
    run.bold = bold
    return p


def build_cover():
    doc = Document()
    style_document(doc)
    doc.add_heading("Cover letter", level=0)
    add_para(doc, "19 September 2026")
    add_para(doc, "Dear Editors,")
    add_para(doc,
        'We submit the manuscript “Preserving physical information in transferable wind-power event measurements” for consideration as a Research Article in Applied Energy.')
    add_para(doc,
        "The manuscript addresses a measurement problem in wind-energy analytics. A dynamic event is created by a detection protocol, matched through a temporal correspondence rule and compressed by a representation. We show why these choices should be evaluated as one measurement interface rather than as interchangeable preprocessing steps.")
    add_para(doc,
        "The study combines seven wind-farm archives and 333 turbines with an independent instrumented French field archive. Matched-event analyses identify a structural reference that persists across detector changes. We then derive a specific information-loss mechanism in signed-domain angular encoding: a power trajectory and its sign reversal share the same angular image. Retaining one polarity coordinate within a six-dimensional representation restores a physically meaningful distinction. Chronological calibration against independent LiDAR measurements raises directional AUROC from 0.494/0.477 to 0.890/0.905 at two turbines. This result connects a representation identity to a measured wind-process outcome.")
    add_para(doc,
        "The paper also tests the operational meaning of the event interface. Three independent observers provide a 320-window human reference, and an observed-price Hill of Towie experiment shows that historical event information reduces two-hour gross imbalance debits by 4.5% in a matched ramp-mixture comparison. Large-ramp periods account for 54–71% of persistence gross debits. These results support a task-aware design principle: event representations should preserve distinctions that survive protocol changes and remain useful for an independent physical or operational task.")
    add_para(doc,
        "The manuscript, supplementary material and source package are prepared from a single canonical version. Public processed data and code are linked to the existing v0.3.0 release; the new v19 ledger and observer records are included as a local release candidate pending final source-licence and archive checks. The study does not present Chinese or French station invoices: the available policy documents define the inputs required for future jurisdiction-specific replay, while the completed monetary result uses Great Britain Elexon prices.")
    add_para(doc,
        "The authors are Peiyan Gong and Chaoxia Yuan of Nanjing University of Information Science and Technology. Funding is provided by the State Key Laboratory of Climate System Prediction and Risk Management initiative project (CPRM-2025-NUIST-012) and the National Natural Science Foundation of China (42088101 and 41875099). The authors declare no competing interests.")
    add_para(doc,
        "The corresponding author is Chaoxia Yuan (chaoxia.yuan@nuist.edu.cn). Before upload, the authors will confirm originality, exclusive submission, all-author approval and data-use permissions. The manuscript contains the required declaration of generative AI and AI-assisted technologies in manuscript preparation.")
    add_para(doc, "Sincerely,")
    add_para(doc, "Peiyan Gong and Chaoxia Yuan")
    doc.save(OUT / "cover_letter_v19.docx")


def build_highlights():
    doc = Document()
    style_document(doc)
    doc.add_heading("Highlights", level=0)
    for text in [
        "Protocol-aware measurement separates event stability from physical information.",
        "Protected polarity restores wind-tendency discrimination after angular encoding.",
        "Independent LiDAR tests connect event representations to measured wind changes.",
        "Historical event information reduces two-hour gross imbalance debits by 4.5%.",
        "Three observers anchor event-presence evaluation across 320 real windows.",
    ]:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(8)
        p.add_run(text)
    doc.save(OUT / "highlights_v19.docx")


if __name__ == "__main__":
    build_cover()
    build_highlights()
