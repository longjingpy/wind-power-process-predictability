"""Create editable journal highlights and an unsigned cover-letter draft."""
from pathlib import Path
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "manuscript/applied_energy"


def document(title):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.85)
    section.left_margin = section.right_margin = Inches(0.9)
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    for name in ["Normal", "Title", "List Bullet"]:
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.size = Pt(12 if name != "Title" else 18)
        style.paragraph_format.space_after = Pt(10)
    doc.add_paragraph(title, "Title")
    return doc


def main():
    meta = json.loads((OUT / "frontmatter.json").read_text(encoding="utf-8"))
    doc = document("Highlights")
    for item in meta["highlights"]:
        assert len(item) <= 85
        doc.add_paragraph(item, "List Bullet")
    doc.save(OUT / "highlights.docx")
    doc = document("Cover letter draft for Applied Energy")
    doc.styles["Normal"].font.size = Pt(10.5)
    doc.styles["Normal"].paragraph_format.space_after = Pt(6)
    doc.styles["Title"].font.size = Pt(16)
    paragraphs = [
        "13 September 2026",
        "Dear Editors,",
        'We propose the manuscript "' + meta["title"] + '" for consideration as a research article in Applied Energy.',
        "Wind-power event catalogues are used to characterize variability, assess forecasting targets and interpret operational responses. These uses assume that the measured event structure survives reasonable changes in detection and representation. Our study tests that assumption before drawing conclusions about event types or weather mechanisms.",
        "The primary comparison covers five wind farms, 189 turbines and 17 detector configurations. Deterministic one-to-one temporal matching separates agreement from the fraction of events that can be matched. The comparison uses identical matched pairs for ordered power shapes, statistical summaries, principal components and image-based representations. Ordered shapes retain substantially more cross-protocol consistency than event-row controls; the tested angular-image pipelines retain less than ordered shapes. These are measurement results, not a claim about every computer-vision model.",
        "A ten-turbine Greek archive provides an external holdout for transfer of the Pizhou-fitted representation. Its native-resolution detector protocol differs from the main physical-time protocol, a boundary stated explicitly. Weather and operational-log analyses are presented as associations and identification diagnostics rather than validated causal effects. The contribution is a practical audit of the event definitions on which downstream energy analyses depend.",
        "The enclosed preparation package contains the manuscript, supplementary methods and result tables, six vector figures and a reproducible typesetting source. The authors are Peiyan Gong and Chaoxia Yuan, affiliated with the State Key Laboratory of Climate System Prediction and Risk Management and the Jiangsu Key Laboratory of Intelligent Weather Forecasting and Applications Based on Big Data, Nanjing University of Information Science and Technology.",
        "This work was financially supported by the State Key Laboratory of Climate System Prediction and Risk Management (CPRM) initiative project (CPRM-2025-NUIST-012) and the National Natural Science Foundation of China (42088101 and 41875099). The authors declare no conflicts of interest.",
        "The Greek monitoring archive and other public-source records will be accompanied by source records and download instructions. Chinese private wind-farm SCADA, coordinates and operational-status records will not be redistributed. The analysis and figure-generation code will be released publicly with documentation and tests; the repository identifier will be added before submission.",
        "The corresponding author is Chaoxia Yuan (chaoxia.yuan@nuist.edu.cn). All-author approval, originality and exclusive-submission status, data-use permissions and the final AI-assistance declaration must be confirmed before upload. This unsigned draft does not assert those confirmations.",
    ]
    for text in paragraphs:
        doc.add_paragraph(text)
    doc.save(OUT / "cover_letter_draft.docx")
    print("Created highlights.docx and cover_letter_draft.docx")


if __name__ == "__main__":
    main()
